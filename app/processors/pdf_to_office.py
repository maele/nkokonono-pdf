from __future__ import annotations

import io
import zipfile
from pathlib import Path

from pypdf import PdfReader

from app.processors.base import BaseProcessor, ProcessorError, ProcessorOutput


class PdfToWordProcessor(BaseProcessor):
    tool_name = "pdf-to-word"

    def validate(self, inputs, options):
        if len(inputs) != 1:
            raise ProcessorError("INVALID_FILE", "PDF to Word accepts exactly one PDF.")
        try:
            import pdf2docx  # noqa: F401
        except ImportError:
            raise ProcessorError("CONVERSION_FAILED", "pdf2docx is not installed.")
        try:
            reader = PdfReader(str(inputs[0].storage_path))
            if reader.is_encrypted:
                raise ProcessorError("ENCRYPTED_PDF", "PDF to Word does not support encrypted PDFs.")
        except ProcessorError:
            raise
        except Exception as exc:
            raise ProcessorError("INVALID_FILE", f"Could not read PDF: {exc}")

    def process(self, inputs, options, output_dir):
        from pdf2docx import Converter

        out = output_dir / "converted.docx"
        try:
            cv = Converter(str(inputs[0].storage_path))
            cv.convert(str(out), start=0, end=None)
            cv.close()
        except Exception as exc:
            raise ProcessorError("CONVERSION_FAILED", f"PDF to Word failed: {exc}")
        if not out.exists():
            raise ProcessorError("OUTPUT_VALIDATION_FAILED", "No output was produced.")
        return [ProcessorOutput(
            path=out,
            display_name="converted.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )]


class PdfToExcelProcessor(BaseProcessor):
    tool_name = "pdf-to-excel"

    def validate(self, inputs, options):
        if len(inputs) != 1:
            raise ProcessorError("INVALID_FILE", "PDF to Excel accepts exactly one PDF.")
        try:
            import pdfplumber  # noqa: F401
            import openpyxl  # noqa: F401
        except ImportError:
            raise ProcessorError("CONVERSION_FAILED", "pdfplumber or openpyxl is not installed.")
        try:
            reader = PdfReader(str(inputs[0].storage_path))
            if reader.is_encrypted:
                raise ProcessorError("ENCRYPTED_PDF", "PDF to Excel does not support encrypted PDFs.")
        except ProcessorError:
            raise
        except Exception as exc:
            raise ProcessorError("INVALID_FILE", f"Could not read PDF: {exc}")

    def process(self, inputs, options, output_dir):
        import openpyxl
        import pdfplumber

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sheet_count = 0
        with pdfplumber.open(str(inputs[0].storage_path)) as pdf:
            for page_idx, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                for table_idx, table in enumerate(tables, start=1):
                    if not table:
                        continue
                    ws = wb.create_sheet(f"p{page_idx}-t{table_idx}")
                    for row in table:
                        ws.append(row)
                    sheet_count += 1
        if sheet_count == 0:
            raise ProcessorError(
                "CONVERSION_FAILED",
                "No tables were detected in this PDF. It may be scanned or use non-tabular layout.",
            )

        out = output_dir / "converted.xlsx"
        try:
            wb.save(str(out))
        except Exception as exc:
            raise ProcessorError("CONVERSION_FAILED", f"Could not write spreadsheet: {exc}")
        return [ProcessorOutput(
            path=out,
            display_name="converted.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )]


class PdfToPptProcessor(BaseProcessor):
    tool_name = "pdf-to-ppt"

    def validate(self, inputs, options):
        if len(inputs) != 1:
            raise ProcessorError("INVALID_FILE", "PDF to PowerPoint accepts exactly one PDF.")
        try:
            import pptx  # noqa: F401
            import fitz  # noqa: F401
        except ImportError:
            raise ProcessorError("CONVERSION_FAILED", "python-pptx or PyMuPDF is not installed.")
        try:
            reader = PdfReader(str(inputs[0].storage_path))
            if reader.is_encrypted:
                raise ProcessorError("ENCRYPTED_PDF", "PDF to PowerPoint does not support encrypted PDFs.")
        except ProcessorError:
            raise
        except Exception as exc:
            raise ProcessorError("INVALID_FILE", f"Could not read PDF: {exc}")

    def process(self, inputs, options, output_dir):
        import fitz
        from pptx import Presentation
        from pptx.util import Inches, Pt

        try:
            dpi = int(options.get("dpi", 150))
        except (TypeError, ValueError):
            raise ProcessorError("INVALID_FILE", "DPI must be an integer.")
        if not (72 <= dpi <= 300):
            raise ProcessorError("INVALID_FILE", "DPI must be between 72 and 300.")

        doc = fitz.open(str(inputs[0].storage_path))
        if doc.page_count == 0:
            doc.close()
            raise ProcessorError("INVALID_FILE", "The PDF has no pages.")

        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 widescreen
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]

        try:
            for i in range(doc.page_count):
                page = doc.load_page(i)
                zoom = dpi / 72.0
                pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                png_bytes = pix.tobytes("png")
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(
                    io.BytesIO(png_bytes),
                    0, 0, width=prs.slide_width, height=prs.slide_height,
                )
        finally:
            doc.close()

        out = output_dir / "converted.pptx"
        try:
            prs.save(str(out))
        except Exception as exc:
            raise ProcessorError("CONVERSION_FAILED", f"Could not write presentation: {exc}")
        return [ProcessorOutput(
            path=out,
            display_name="converted.pptx",
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )]